"""Build the Word manuscript (paper/manuscript.docx) from paper/manuscript.md.

Supported markup: '# ' title, '## ' / '### ' headings, paragraphs with **bold** runs, '- ' bullets,
'1. ' numbered references, pipe tables (preceded by a '**Table N.** caption' paragraph),
and figures as '![caption](path)'. Layout follows Elsevier submission norms: A4, 2.5 cm margins,
Times New Roman 12 pt, double spacing, continuous line numbers, page numbers in the footer.

Usage: .venv/bin/python -I paper/build_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FONT = "Times New Roman"


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_runs(par, text, size=None, italic=False):
    """Add text to a paragraph, rendering **bold** spans."""
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not chunk:
            continue
        r = par.add_run(chunk)
        r.bold = i % 2 == 1
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    return par


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def line_numbers(section):
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    section._sectPr.append(ln)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\|\s*-", r)]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j, text in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, text, size=9)
            if i == 0:
                for r in p.runs:
                    r.bold = True
                set_cell_shading(c, "E8E8E4")
    doc.add_paragraph()


def main():
    md = (HERE / "manuscript.md").read_text().splitlines()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    st.paragraph_format.space_after = Pt(0)
    for name, size in [("Title", 16), ("Heading 1", 13), ("Heading 2", 12)]:
        s = doc.styles[name]
        s.font.name = FONT
        s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    line_numbers(sec)
    page_number_footer(sec)

    i, in_refs = 0, False
    while i < len(md):
        line = md[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(12)
            r = p.add_run(line[2:])
            r.bold, r.font.size = True, Pt(16)
        elif line.startswith("## "):
            in_refs = line[3:].strip() == "References"
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.*)\)", line)
            cap, path = m.group(1), HERE / m.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Cm(16))
            label, _, rest = cap.partition(". ")
            c = doc.add_paragraph()
            c.paragraph_format.line_spacing = 1.15
            c.paragraph_format.space_after = Pt(12)
            r = c.add_run(label + ". ")
            r.bold = True
            r.font.size = Pt(10)
            add_runs(c, rest, size=10)
        elif line.startswith("|"):
            rows = []
            while i < len(md) and md[i].startswith("|"):
                rows.append(md[i])
                i += 1
            table(doc, rows)
            continue
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
        elif in_refs and re.match(r"^\d+\. ", line):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.8)
            p.paragraph_format.space_after = Pt(4)
            num, _, rest = line.partition(". ")
            add_runs(p, f"[{num}] {rest}", size=10.5)
        elif line.startswith("**Table"):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(6)
            add_runs(p, line, size=10)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0) if line.startswith(("**", "[", "Corresponding")) else Cm(0.6)
            add_runs(p, line)
        i += 1
    out = HERE / "manuscript.docx"
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
